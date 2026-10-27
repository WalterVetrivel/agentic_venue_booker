"""Idempotent seed: roles, users (from the faculty Excel sheet), venues and equipment."""

import re
from pathlib import Path

import yaml
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.auth import hash_password
from app.config import get_settings
from app.models import Equipment, Role, SeasonalRule, User, UserRole, Venue, VenueAlias
from app.permissions import ROLES

SEED_DIR = Path(__file__).resolve().parent
DEFAULT_PASSWORD = "password"

# People whose role differs from the sheet (agreed with Walter).
DIRECTOR = {"name": "Dr. V. Kumar", "email": "dirim@kiot.ac.in", "designation": "Director - IM"}
HOD = {"name": "Mr. T. Karthikeyan", "email": "hod.cse@kiot.ac.in", "designation": "HOD / CSE"}
OA = {"name": "CSE Office Assistant", "email": "cseoa@kiot.ac.in", "designation": "OA / DEO"}
NO_FACULTY_LOGIN = {"VK"}  # Dr. V. Kumar only gets the Director login


def normalise_name(name: str) -> str:
    name = re.sub(r"\.(?=\S)", ". ", name.strip())
    return re.sub(r"\s+", " ", name)


def faculty_email(initials: str) -> str:
    return f"{initials.lower()}cse@kiot.ac.in"


def read_faculty(xlsx: Path) -> list[dict]:
    ws = load_workbook(xlsx, read_only=True).active
    rows = ws.iter_rows(min_row=2, values_only=True)
    faculty = []
    for _sno, name, designation, initials, *_ in rows:
        if not name or not initials:
            continue
        faculty.append(
            {
                "name": normalise_name(str(name)),
                "designation": str(designation or "").strip(),
                "initials": str(initials).strip().upper(),
            }
        )
    return faculty


def _upsert_user(db: Session, pw_hash: str, *, email: str, roles: list[str], **fields) -> User:
    user = db.scalar(select(User).where(User.email == email))
    if user is None:
        user = User(email=email, password_hash=pw_hash, **fields)
        db.add(user)
    have = {r.role_key for r in user.roles}
    for key in roles:
        if key not in have:
            user.roles.append(UserRole(role_key=key))
    return user


def seed_all(db: Session) -> dict[str, int]:
    settings = get_settings()
    for key, label in ROLES.items():
        if db.get(Role, key) is None:
            db.add(Role(key=key, label=label))
    db.flush()

    pw_hash = hash_password(DEFAULT_PASSWORD)  # hashed once; all mock accounts share it

    for f in read_faculty(settings.mock_data_dir / "faculty.xlsx"):
        if f["initials"] in NO_FACULTY_LOGIN:
            continue
        _upsert_user(db, pw_hash, email=faculty_email(f["initials"]), roles=["faculty"], **f)
    _upsert_user(db, pw_hash, roles=["director"], **DIRECTOR)
    _upsert_user(db, pw_hash, roles=["hod"], initials="TK", **HOD)
    _upsert_user(db, pw_hash, roles=["oa"], **OA)

    data = yaml.safe_load((SEED_DIR / "venues.yaml").read_text(encoding="utf-8"))
    for h in data["hall_incharges"]:
        _upsert_user(db, pw_hash, roles=["incharge"], designation="Hall In-charge", **h)
    db.flush()

    for v in data["venues"]:
        _seed_venue(db, pw_hash, v, data["software_profiles"])
    db.commit()

    return {
        "users": len(db.scalars(select(User)).all()),
        "venues": len(db.scalars(select(Venue)).all()),
        "equipment": len(db.scalars(select(Equipment)).all()),
    }


def _seed_venue(db: Session, pw_hash: str, v: dict, profiles: dict[str, list[str]]) -> None:
    if db.scalar(select(Venue).where(Venue.code == v["code"])) is not None:
        return
    in_charge = db.scalar(select(User).where(User.email == v["in_charge"]))
    if in_charge is not None and "incharge" not in in_charge.role_keys:
        in_charge.roles.append(UserRole(role_key="incharge"))

    is_cse = v.get("department", "CSE") == "CSE"
    venue = Venue(
        code=v["code"],
        name=v["name"],
        type=v["type"],
        block=v.get("block"),
        department=v.get("department", "CSE"),
        capacity=v["capacity"],
        byod=v.get("byod", False),
        needs_hod_approval=is_cse,
        in_charge=in_charge,
    )
    db.add(venue)

    aliases = {v["code"], v["name"], *v.get("aliases", [])}
    if v["type"] == "lab":
        aliases.add(f"{v['code']} Lab")
    if v["type"] == "coe":
        n = v["code"].removeprefix("COE")
        aliases |= {
            f"{v['code']} Lab",
            f"Centre of Excellence ({v['code']})",
            f"Centre of Excellence {n}",
            f"CoE {n}",
        }
    venue.aliases = [VenueAlias(alias=a) for a in sorted(aliases)]

    items: list[Equipment] = []
    if v["type"] == "seminar_hall":
        av = [
            "Projector",
            "PA system",
            "Podium microphone",
            "Cordless microphones",
            "Air conditioning",
        ]
    else:
        av = ["Projector", "Smart board", "Air conditioning"]
    items += [Equipment(name=n, category="av") for n in av]
    if v.get("pcs"):
        items.append(
            Equipment(
                name="Desktop PC (Windows 11, i5, 16 GB RAM)",
                category="hardware",
                quantity=v["pcs"],
            )
        )
    for hw in v.get("extra_hardware", []):
        name, _, qty = hw.rpartition(" x")
        items.append(
            Equipment(
                name=name or hw, category="hardware", quantity=int(qty) if qty.isdigit() else 1
            )
        )
    for sw in profiles.get(v.get("software", ""), []):
        items.append(Equipment(name=sw, category="software", quantity=v.get("pcs", 1)))
    if v.get("byod"):
        items += [
            Equipment(name="Power points for laptops", category="facility", quantity=30),
            Equipment(name="Wi-Fi", category="facility"),
        ]
    venue.equipment = items

    venue.seasonal_rules = [SeasonalRule(**r) for r in v.get("seasonal_rules", [])]
